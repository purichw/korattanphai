#!/usr/bin/env python3
from __future__ import annotations

import argparse
import json
import re
from collections import Counter, defaultdict
from datetime import datetime, timezone, timedelta
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
ADMIN_JSON = ROOT / "src/data/canonical/nakhon_ratchasima/admin_hierarchy.json"
GEODATA_JSON = ROOT / "public/geodata/nakhon-ratchasima-subdistricts.geojson"
DEFAULT_OUTPUT = ROOT / "src/data/canonical/nakhon_ratchasima/drought_forecast_archive_rev02.json"

PROVINCE_ID = "TH-P29"
PROVINCE_CODE = "30"
PROVINCE_TH = "นครราชสีมา"
PROVINCE_EN = "Nakhon Ratchasima"
SOURCE_WORKBOOK_ORIGINAL = "Drought_T1-6_rev02.xlsx"

DISTRICT_ALIASES = {
    "CHALOEMPHRAKIET": "CHALEOMPHRAKIAT",
    "KHAMSAKAESAENG": "KHAMSAKAESAENG",
    "LAMTAMANCHAI": "LAMTHAMENCHAI",
    "SIKHIO": "SIKHIO",
}

GEODATA_SUBDISTRICT_EN_OVERRIDES = {
    "300101": "TAMBON NAI MUEANG",
    "300113": "TAMBON BAN MAI",
    "301811": "TAMBON KUT CHIK",
    "302401": "TAMBON NON DAENG",
    "302405": "TAMBON DON YAO YAI",
}

TH_MONTH_ABBR = ["ม.ค.", "ก.พ.", "มี.ค.", "เม.ย.", "พ.ค.", "มิ.ย.", "ก.ค.", "ส.ค.", "ก.ย.", "ต.ค.", "พ.ย.", "ธ.ค."]
RISK_LABEL_TH = {
    0: "ไม่มีความเสี่ยง",
    1: "เสี่ยงปานกลาง",
    2: "เสี่ยงสูง",
    None: "นอกขอบเขตการศึกษา",
}


def compact_name(value: object) -> str:
    text = str(value or "").upper().strip()
    text = re.sub(r"^(AMPHOE|TAMBON|CHANGWAT)\s+", "", text)
    text = re.sub(r"[^A-Z0-9]+", " ", text)
    text = re.sub(r"\s+", " ", text).strip().replace(" ", "")
    return DISTRICT_ALIASES.get(text, text)


def normalize_irrigation_status(value: object) -> tuple[str, str]:
    raw = str(value or "").strip()
    if raw == "Irragation":
        return "Irrigation", "IRRIGATION_STATUS_NORMALIZED"
    return raw, ""


def read_records(workbook: Any, sheet_name: str) -> list[dict[str, Any]]:
    worksheet = workbook[sheet_name]
    rows = worksheet.iter_rows(values_only=True)
    headers = [str(value).strip() for value in next(rows)]
    records = []
    for row in rows:
        if row is None or not any(value is not None for value in row):
            continue
        records.append(dict(zip(headers, row)))
    return records


def month_shift(period: str, offset: int) -> str:
    year_text, month_text = period.split("-")
    total_month = int(year_text) * 12 + int(month_text) - 1 + offset
    return f"{total_month // 12}-{(total_month % 12) + 1:02d}"


def label_month_th(period: str) -> str:
    year_text, month_text = period.split("-")
    return f"{TH_MONTH_ABBR[int(month_text) - 1]} {int(year_text) + 543}"


def load_canonical_geography() -> dict[str, list[dict[str, Any]]]:
    admin = json.loads(ADMIN_JSON.read_text(encoding="utf-8"))
    canonical_by_code: dict[str, dict[str, Any]] = {}
    for district in admin["province"]["districts"]:
        for subdistrict in district["subdistricts"]:
            canonical_by_code[subdistrict["subdistrictCode"]] = {
                "provinceId": PROVINCE_ID,
                "provinceCode": PROVINCE_CODE,
                "provinceNameTh": PROVINCE_TH,
                "provinceNameEn": PROVINCE_EN,
                "districtCode": district["districtCode"],
                "districtId": district["id"],
                "districtNameTh": district["nameTh"],
                "districtNameEn": district["name"],
                "districtSlug": district["routingSlug"],
                "subdistrictCode": subdistrict["subdistrictCode"],
                "subdistrictId": subdistrict["id"],
                "subdistrictNameTh": subdistrict["nameTh"],
                "subdistrictSlug": subdistrict["routingSlug"],
            }

    geodata = json.loads(GEODATA_JSON.read_text(encoding="utf-8"))
    by_join_key: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for feature in geodata["features"]:
        props = feature["properties"]
        subdistrict_code = str(props["Admin_code"])
        canonical = canonical_by_code[subdistrict_code]
        geodata_subdistrict_name_for_match = GEODATA_SUBDISTRICT_EN_OVERRIDES.get(
            subdistrict_code,
            props["T_Name_E"],
        )
        join_key = f"{compact_name(props['A_Name_E'])}|{compact_name(geodata_subdistrict_name_for_match)}"
        by_join_key[join_key].append(
            {
                **canonical,
                "geodataDistrictNameEn": props["A_Name_E"],
                "geodataSubdistrictNameEnOriginal": props["T_Name_E"],
                "geodataSubdistrictNameEnForMatch": geodata_subdistrict_name_for_match,
                "geodataLabelOverrideApplied": subdistrict_code in GEODATA_SUBDISTRICT_EN_OVERRIDES,
            }
        )
    return by_join_key


def build_location_mapping(location_master: list[dict[str, Any]]) -> tuple[list[dict[str, Any]], dict[str, dict[str, Any]], dict[str, Any]]:
    geography_by_join_key = load_canonical_geography()
    source_id_counts = Counter(str(int(row["Source_ID"])) for row in location_master)
    duplicate_source_ids = sorted(source_id for source_id, count in source_id_counts.items() if count > 1)
    mapped_locations: list[dict[str, Any]] = []
    mapped_by_source_id: dict[str, dict[str, Any]] = {}
    unmatched: list[dict[str, Any]] = []
    ambiguous: list[dict[str, Any]] = []

    for row in location_master:
        source_id = str(int(row["Source_ID"]))
        source_tambon_en = str(row["TAMBON_E"]).strip()
        source_amphoe_en = str(row["AMPHOE_E"]).strip()
        corrected_amphoe_en = "Phimai" if source_id == "222" and source_tambon_en == "Nong Rawiang" else source_amphoe_en
        source_admin_correction_applied = corrected_amphoe_en != source_amphoe_en
        irrigation_status, irrigation_quality_flag = normalize_irrigation_status(row["Irrigation_Status"])
        join_key = f"{compact_name(corrected_amphoe_en)}|{compact_name(source_tambon_en)}"
        matches = geography_by_join_key.get(join_key, [])

        if not matches:
            unmatched.append({"sourceId": source_id, "sourceAmphoeEn": source_amphoe_en, "sourceTambonEn": source_tambon_en})
            continue
        if len(matches) > 1:
            ambiguous.append(
                {
                    "sourceId": source_id,
                    "sourceAmphoeEn": source_amphoe_en,
                    "sourceTambonEn": source_tambon_en,
                    "candidateSubdistrictCodes": [match["subdistrictCode"] for match in matches],
                }
            )
            continue

        match = matches[0]
        methods = []
        if source_admin_correction_applied:
            methods.append("SOURCE_ADMIN_CORRECTION")
        if match["geodataLabelOverrideApplied"]:
            methods.append("GEODATA_EN_LABEL_OVERRIDE")
        if compact_name(source_amphoe_en) == compact_name(match["geodataDistrictNameEn"]) and compact_name(source_tambon_en) == compact_name(
            match["geodataSubdistrictNameEnOriginal"]
        ):
            methods.append("EXACT_NORMALIZED_EN")
        else:
            methods.append("ALIAS_NORMALIZED_EN")
        if irrigation_quality_flag:
            methods.append(irrigation_quality_flag)

        quality_flags = []
        if source_admin_correction_applied:
            quality_flags.append("SOURCE_AMPHOE_CORRECTED")
        if match["geodataLabelOverrideApplied"]:
            quality_flags.append("GEODATA_LABEL_OVERRIDE")
        if irrigation_quality_flag:
            quality_flags.append(irrigation_quality_flag)

        mapped = {
            "sourceId": source_id,
            "sourceAreaKey": f"{source_amphoe_en} / {source_tambon_en} / ID {source_id}",
            "sourceTambonEn": source_tambon_en,
            "sourceAmphoeEn": source_amphoe_en,
            "sourceAmphoeEnCorrected": corrected_amphoe_en,
            "sourceAdminCorrectionApplied": source_admin_correction_applied,
            "irrigationStatusRaw": str(row["Irrigation_Status"]).strip(),
            "irrigationStatus": irrigation_status,
            "provinceId": match["provinceId"],
            "provinceCode": match["provinceCode"],
            "provinceNameTh": match["provinceNameTh"],
            "provinceNameEn": match["provinceNameEn"],
            "districtCode": match["districtCode"],
            "districtId": match["districtId"],
            "districtNameTh": match["districtNameTh"],
            "districtNameEn": match["districtNameEn"],
            "districtSlug": match["districtSlug"],
            "subdistrictCode": match["subdistrictCode"],
            "subdistrictId": match["subdistrictId"],
            "subdistrictNameTh": match["subdistrictNameTh"],
            "subdistrictSlug": match["subdistrictSlug"],
            "matchStatus": "MATCHED",
            "matchMethod": ";".join(dict.fromkeys(methods)),
            "qualityFlags": ";".join(quality_flags),
        }
        mapped_locations.append(mapped)
        mapped_by_source_id[source_id] = mapped

    canonical_counts = Counter(location["subdistrictCode"] for location in mapped_locations)
    duplicate_canonical = sorted(code for code, count in canonical_counts.items() if count > 1)

    diagnostics = {
        "sourceIdCount": len(source_id_counts),
        "mappedSourceIdCount": len(mapped_by_source_id),
        "mappedCanonicalSubdistrictCount": len(canonical_counts),
        "duplicateSourceIds": duplicate_source_ids,
        "unmappedSourceIds": [item["sourceId"] for item in unmatched],
        "unmatchedSources": unmatched,
        "ambiguousSourceIds": [item["sourceId"] for item in ambiguous],
        "ambiguousSources": ambiguous,
        "duplicateCanonicalSubdistrictCodes": duplicate_canonical,
        "sourceCorrections": [
            {
                "sourceId": location["sourceId"],
                "sourceAmphoeEn": location["sourceAmphoeEn"],
                "sourceAmphoeEnCorrected": location["sourceAmphoeEnCorrected"],
                "sourceTambonEn": location["sourceTambonEn"],
                "subdistrictCode": location["subdistrictCode"],
            }
            for location in mapped_locations
            if location["sourceAdminCorrectionApplied"]
        ],
        "geodataLabelOverrides": [
            {"subdistrictCode": code, "override": label}
            for code, label in sorted(GEODATA_SUBDISTRICT_EN_OVERRIDES.items())
        ],
    }
    if duplicate_source_ids or unmatched or ambiguous or duplicate_canonical:
        raise RuntimeError(json.dumps(diagnostics, ensure_ascii=False, indent=2))

    return (
        sorted(mapped_locations, key=lambda item: (item["districtCode"], item["subdistrictCode"])),
        mapped_by_source_id,
        diagnostics,
    )


def read_data_quality(records: list[dict[str, Any]]) -> dict[str, Any]:
    metrics: dict[str, Any] = {}
    metric_keys = {
        "Source rows": "sourceRowCountOriginal",
        "Exact unique rows": "sourceRowCountDeduped",
        "Exact duplicates removed": "duplicateSourceRowsRemoved",
        "Unique Source_IDs": "sourceIdCount",
        "Unique Source_YearMonth": "sourceYearMonthCount",
        "Period start": "periodStart",
        "Period end": "periodEnd",
        "Long rows": "longRowCount",
    }
    for row in records:
        key = metric_keys.get(str(row.get("Metric", "")).strip())
        if key:
            metrics[key] = row.get("Value")
    return metrics


def parse_risk(value: Any) -> int | None:
    if value is None or value == "":
        return None
    numeric = int(value)
    if numeric not in (0, 1, 2):
        raise ValueError(f"Invalid forecast risk: {value}")
    return numeric


def build_archive(
    workbook_path: Path,
    output_path: Path,
) -> dict[str, Any]:
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    location_master = read_records(workbook, "Location_Master")
    long_records = read_records(workbook, "Forecast_Archive_Long")
    data_quality = read_data_quality(read_records(workbook, "Data_Quality"))
    locations, location_by_source_id, mapping_diagnostics = build_location_mapping(location_master)
    location_by_subdistrict_code = {location["subdistrictCode"]: location for location in locations}

    packed_risk_by_target_month: dict[str, dict[str, list[int | None]]] = defaultdict(dict)
    source_vintage_keys: set[str] = set()
    canonical_vintage_keys: set[str] = set()
    period_horizon_counts: dict[tuple[str, int], Counter] = defaultdict(Counter)
    horizon_counts: dict[int, Counter] = defaultdict(Counter)
    target_months: set[str] = set()
    issue_months: set[str] = set()
    source_period_counts = Counter()
    equal_valued_example: dict[str, Any] | None = None

    for row in long_records:
        source_id = str(int(row["Source_ID"]))
        location = location_by_source_id[source_id]
        subdistrict_code = location["subdistrictCode"]
        target_month = str(row["Source_YearMonth"])
        horizon = int(row["Horizon"])
        if horizon not in (1, 2, 3, 4, 5, 6):
            raise ValueError(f"Invalid horizon: {horizon}")
        forecast_risk = parse_risk(row["Forecast_Risk"])
        workbook_scope = str(row["Scope_Status"])
        scope_status = "out_of_scope" if forecast_risk is None else "in_scope"
        if workbook_scope != scope_status:
            raise ValueError(
                f"Scope mismatch for Source_ID {source_id}, {target_month}, T+{horizon}: "
                f"{workbook_scope} vs {scope_status}"
            )

        issue_month = month_shift(target_month, -horizon)
        target_months.add(target_month)
        issue_months.add(issue_month)
        source_period_counts[target_month] += 1
        source_vintage_keys.add(str(row["Vintage_Key_Source"]))
        canonical_vintage_keys.add(f"{subdistrict_code}|{target_month}|T+{horizon}")
        target_bucket = packed_risk_by_target_month[target_month]
        risks = target_bucket.setdefault(subdistrict_code, [None, None, None, None, None, None])
        risks[horizon - 1] = forecast_risk

        scope_key = "outOfScope" if forecast_risk is None else "inScope"
        risk_key = "out_of_scope" if forecast_risk is None else str(forecast_risk)
        period_horizon_counts[(target_month, horizon)][scope_key] += 1
        period_horizon_counts[(target_month, horizon)][risk_key] += 1
        horizon_counts[horizon][scope_key] += 1
        horizon_counts[horizon][risk_key] += 1

    expected_source_ids = len(location_by_source_id)
    for period, subdistrict_map in packed_risk_by_target_month.items():
        if len(subdistrict_map) != expected_source_ids:
            raise RuntimeError(f"{period} has {len(subdistrict_map)} subdistricts, expected {expected_source_ids}")
        for subdistrict_code, risks in subdistrict_map.items():
            if len(risks) != 6:
                raise RuntimeError(f"{period}/{subdistrict_code} has {len(risks)} horizons")
            source_id = location_by_subdistrict_code[subdistrict_code]["sourceId"]
            if source_id == "101" and period == "2025-12":
                equal_valued_example = {
                    "subdistrictCode": subdistrict_code,
                    "sourceId": source_id,
                    "targetMonth": period,
                    "risksByHorizon": risks,
                    "vintageKeys": [f"{subdistrict_code}|{period}|T+{horizon}" for horizon in range(1, 7)],
                    "issueMonths": [month_shift(period, -horizon) for horizon in range(1, 7)],
                }

    target_month_list = sorted(target_months)
    target_month_records = []
    for period in target_month_list:
        horizon_summaries = []
        for horizon in range(1, 7):
            counts = period_horizon_counts[(period, horizon)]
            horizon_summaries.append(
                {
                    "horizon": horizon,
                    "horizonLabel": f"T+{horizon}",
                    "issueMonth": month_shift(period, -horizon),
                    "targetMonth": period,
                    "vintageCount": counts["inScope"] + counts["outOfScope"],
                    "inScopeSubdistricts": counts["inScope"],
                    "outOfScopeSubdistricts": counts["outOfScope"],
                    "noRiskSubdistricts": counts["0"],
                    "moderateRiskSubdistricts": counts["1"],
                    "highRiskSubdistricts": counts["2"],
                }
            )
        target_month_records.append(
            {
                "period": period,
                "labelTh": label_month_th(period),
                "horizons": horizon_summaries,
            }
        )

    horizon_summary = []
    for horizon in range(1, 7):
        counts = horizon_counts[horizon]
        horizon_summary.append(
            {
                "horizon": horizon,
                "horizonLabel": f"T+{horizon}",
                "vintageCount": counts["inScope"] + counts["outOfScope"],
                "inScopeVintages": counts["inScope"],
                "outOfScopeVintages": counts["outOfScope"],
                "noRiskVintages": counts["0"],
                "moderateRiskVintages": counts["1"],
                "highRiskVintages": counts["2"],
            }
        )

    latest_period = target_month_list[-1]
    dan_khun_thot_codes = [
        location["subdistrictCode"] for location in locations if location["districtCode"] == "3008"
    ]
    latest_t1_counts = period_horizon_counts[(latest_period, 1)]
    dan_khun_thot_t1 = Counter()
    for code in dan_khun_thot_codes:
        risk = packed_risk_by_target_month[latest_period][code][0]
        if risk is None:
            dan_khun_thot_t1["outOfScope"] += 1
        else:
            dan_khun_thot_t1["inScope"] += 1
            dan_khun_thot_t1[str(risk)] += 1

    now = datetime.now(timezone(timedelta(hours=7))).isoformat(timespec="seconds")
    archive = {
        "meta": {
            "sourceOfTruth": "normalized_rev02_forecast_archive_workbook",
            "sourceWorkbook": workbook_path.name,
            "sourceWorkbookOriginal": SOURCE_WORKBOOK_ORIGINAL,
            "sourceSheet": "Forecast_Archive_Long",
            "locationSheet": "Location_Master",
            "generatedAt": now,
            "timezone": "Asia/Bangkok",
            "provinceId": PROVINCE_ID,
            "provinceCode": PROVINCE_CODE,
            "provinceNameTh": PROVINCE_TH,
            "provinceNameEn": PROVINCE_EN,
            "provenance": "REAL",
            "semanticsTh": "คลังคำพยากรณ์ภัยแล้ง T+1 ถึง T+6 จาก workbook rev02; blank หมายถึงนอกขอบเขตการศึกษา ไม่ใช่ไม่มีความเสี่ยง",
            "sourceRowCountOriginal": int(data_quality.get("sourceRowCountOriginal", 0)),
            "sourceRowCountDeduped": int(data_quality.get("sourceRowCountDeduped", 0)),
            "duplicateSourceRowsRemoved": int(data_quality.get("duplicateSourceRowsRemoved", 0)),
            "sourceIdCount": len(location_by_source_id),
            "targetMonthCount": len(target_month_list),
            "horizonCount": 6,
            "forecastVintageCount": len(canonical_vintage_keys),
            "sourceVintageKeyCount": len(source_vintage_keys),
            "forecastVintageIdentity": "subdistrictCode + targetMonth + horizon",
            "totalCanonicalSubdistricts": len(location_by_subdistrict_code),
            "periodStart": target_month_list[0],
            "periodEnd": latest_period,
            "targetMonthStart": target_month_list[0],
            "targetMonthEnd": latest_period,
            "issueMonthStart": sorted(issue_months)[0],
            "issueMonthEnd": sorted(issue_months)[-1],
            "temporalInterpretation": "SOURCE_YEARMONTH_IS_TARGET_MONTH",
            "targetMonthRule": "targetMonth = Source_YearMonth",
            "issueMonthRule": "issueMonth = subtractMonths(targetMonth, horizon)",
            "temporalInterpretationEvidence": [
                "Existing archive route state already uses the query parameter name target for selected map month.",
                "Previous product examples for Ban Kao / Dan Khun Thot match Source_YearMonth 2025-12 as target month with T+1 issued in 2025-11 and T+2 issued in 2025-10.",
                "The workbook field name stays preserved as sourceYearMonth in lineage while the product-facing archive uses targetMonth.",
            ],
        },
        "riskSemantics": [
            {"forecastRisk": 0, "scopeStatus": "in_scope", "labelTh": RISK_LABEL_TH[0], "mapStatus": "forecast-no-risk"},
            {"forecastRisk": 1, "scopeStatus": "in_scope", "labelTh": RISK_LABEL_TH[1], "mapStatus": "forecast-moderate"},
            {"forecastRisk": 2, "scopeStatus": "in_scope", "labelTh": RISK_LABEL_TH[2], "mapStatus": "forecast-high"},
            {
                "forecastRisk": None,
                "scopeStatus": "out_of_scope",
                "labelTh": RISK_LABEL_TH[None],
                "mapStatus": "forecast-out-of-scope",
            },
        ],
        "mapping": mapping_diagnostics,
        "locations": locations,
        "targetMonths": target_month_records,
        "horizonSummary": horizon_summary,
        "validationExamples": {
            "equalValuedTPlusVintagesRemainSeparate": equal_valued_example,
            "canonicalSubdistrict300806": {
                "sourceId": "101",
                "districtCode": "3008",
                "districtNameTh": "ด่านขุนทด",
                "subdistrictCode": "300806",
                "subdistrictNameTh": "บ้านเก่า",
                "targetMonth": latest_period,
                "risksByHorizon": packed_risk_by_target_month[latest_period]["300806"],
                "issueMonths": [month_shift(latest_period, -horizon) for horizon in range(1, 7)],
            },
            "provinceLatestT1": {
                "targetMonth": latest_period,
                "horizon": 1,
                "issueMonth": month_shift(latest_period, -1),
                "totalSubdistricts": len(location_by_subdistrict_code),
                "inScopeSubdistricts": latest_t1_counts["inScope"],
                "noRiskSubdistricts": latest_t1_counts["0"],
                "moderateRiskSubdistricts": latest_t1_counts["1"],
                "highRiskSubdistricts": latest_t1_counts["2"],
                "outOfScopeSubdistricts": latest_t1_counts["outOfScope"],
            },
            "danKhunThotLatestT1": {
                "districtCode": "3008",
                "districtNameTh": "ด่านขุนทด",
                "targetMonth": latest_period,
                "horizon": 1,
                "issueMonth": month_shift(latest_period, -1),
                "totalSubdistricts": len(dan_khun_thot_codes),
                "inScopeSubdistricts": dan_khun_thot_t1["inScope"],
                "noRiskSubdistricts": dan_khun_thot_t1["0"],
                "moderateRiskSubdistricts": dan_khun_thot_t1["1"],
                "highRiskSubdistricts": dan_khun_thot_t1["2"],
                "outOfScopeSubdistricts": dan_khun_thot_t1["outOfScope"],
            },
        },
        "packedRiskByTargetMonth": {
            period: dict(sorted(subdistrict_map.items()))
            for period, subdistrict_map in sorted(packed_risk_by_target_month.items())
        },
    }

    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(json.dumps(archive, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return archive


def main() -> None:
    parser = argparse.ArgumentParser(description="Build the Nakhon Ratchasima drought forecast archive fixture.")
    parser.add_argument("workbook", type=Path, help="Path to Drought_T1-6_rev02_Normalized_ArchiveReady.xlsx")
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT, help="Canonical JSON output path")
    args = parser.parse_args()

    archive = build_archive(args.workbook, args.output)
    summary = {
        "output": str(args.output),
        "sourceIdCount": archive["meta"]["sourceIdCount"],
        "mappedCanonicalSubdistrictCount": archive["mapping"]["mappedCanonicalSubdistrictCount"],
        "targetMonthStart": archive["meta"]["targetMonthStart"],
        "targetMonthEnd": archive["meta"]["targetMonthEnd"],
        "issueMonthStart": archive["meta"]["issueMonthStart"],
        "issueMonthEnd": archive["meta"]["issueMonthEnd"],
        "forecastVintageCount": archive["meta"]["forecastVintageCount"],
    }
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
