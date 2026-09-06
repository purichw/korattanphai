#!/usr/bin/env python3
"""Normalize rev03 using only the workbook and the user's forward-month definitions."""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook

SHEET = "Master_Data_Drought_Final"
HEADERS = ["Year", "Month", "ID", "TAMBON_E", "AMPHOE_E", "Irrigation_Status"] + [f"Pred_Risk_T+{h}" for h in range(1, 7)]
RISK_LABELS = {"0": "ไม่พบสัญญาณเสี่ยง", "1": "เสี่ยงปานกลาง", "2": "เสี่ยงสูง", "null": "อยู่นอกขอบเขตการศึกษา (การพยากรณ์)"}
IRRIGATION_LABELS = {"Collecting": "ยังไม่มีข้อมูล", "RainFed": "พึ่งน้ำฝน (ไม่มีชลประทาน)", "Irrigation": "เข้าถึงชลประทาน"}


def sha256(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def integer(value, name, minimum, maximum):
    if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value) or value != int(value) or not minimum <= value <= maximum:
        raise ValueError(f"Invalid {name}: {value!r}")
    return int(value)


def risk_value(value):
    if value is None or value == "":
        return None
    return integer(value, "risk", 0, 2)


def irrigation_status(value):
    normalized = "Irrigation" if value == "Irragation" else value
    if normalized not in IRRIGATION_LABELS:
        raise ValueError(f"Invalid irrigation status: {value!r}")
    return normalized


def target_month(source_month, horizon):
    integer(horizon, "horizon", 1, 6)
    year, month = map(int, source_month.split("-"))
    integer(year, "year", 1900, 2199)
    integer(month, "month", 1, 12)
    target = year * 12 + month - 1 + horizon
    return f"{target // 12:04d}-{target % 12 + 1:02d}"


def normalize_rows(numbered_rows):
    unique, locations, lineage = {}, {}, []
    blank_rows, risk_types, raw_counts = [], Counter(), Counter()
    for number, raw in numbered_rows:
        raw = tuple(raw)
        if all(value is None for value in raw):
            blank_rows.append(number)
            continue
        if len(raw) != len(HEADERS):
            raise ValueError(f"Row {number}: unexpected column count")
        year = integer(raw[0], "Gregorian year", 1900, 2199)
        month = integer(raw[1], "month", 1, 12)
        source_id = str(integer(raw[2], "source ID", 0, 999999))
        if not all(isinstance(value, str) and value.strip() for value in raw[3:6]):
            raise ValueError(f"Row {number}: missing location metadata")
        irrigation_status(raw[5])
        risks = [risk_value(value) for value in raw[6:]]
        risk_types.update(type(value).__name__ for value in raw[6:])
        raw_counts.update("null" if value is None else str(value) for value in risks)
        source_month = f"{year:04d}-{month:02d}"
        key = (source_month, source_id)
        if source_id in locations and locations[source_id] != raw[3:6]:
            raise ValueError(f"Row {number}: changing location/irrigation metadata for ID {source_id}")
        locations[source_id] = raw[3:6]
        if key in unique:
            if unique[key]["raw"] != raw:
                raise ValueError(f"Row {number}: conflicting duplicate for {key}; no last-row-wins")
            unique[key]["source_row_count"] += 1
        else:
            unique[key] = {"raw": raw, "risks": risks, "source_first_row": number, "source_row_count": 1}
        lineage.append({"source_row": number, "source_id": source_id, "source_month": source_month,
                        "retained_source_row": unique[key]["source_first_row"]})
    if not unique:
        raise ValueError("No forecast rows")
    months = sorted({key[0] for key in unique})
    expected_months = (int(months[-1][:4]) - int(months[0][:4])) * 12 + int(months[-1][5:]) - int(months[0][5:]) + 1
    if len(months) != expected_months or len(unique) != len(months) * len(locations):
        raise ValueError("Incomplete location/month grid; do not create blank forecasts for missing rows")
    quality = {"source_rows": len(lineage), "exact_unique_rows": len(unique),
               "exact_duplicates_removed": len(lineage) - len(unique), "blank_row_numbers": blank_rows,
               "conflicting_keys": 0, "location_metadata_conflicts": 0,
               "source_month_count": len(months), "source_month_start": months[0], "source_month_end": months[-1],
               "source_id_count": len(locations), "raw_risk_cell_types": dict(risk_types),
               "raw_risk_counts_including_duplicates": dict(raw_counts),
               "duplicate_multiplicity": dict(sorted(Counter(row["source_row_count"] for row in unique.values()).items()))}
    return unique, locations, lineage, quality


def write_json(path, value):
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def write_csv(path, rows, fields):
    with path.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=fields, lineterminator="\n")
        writer.writeheader()
        writer.writerows(rows)


def build(workbook_path, output):
    if output.exists():
        raise ValueError(f"Output already exists; choose a new directory: {output}")
    source_hash = sha256(workbook_path)
    workbook = load_workbook(workbook_path, read_only=True, data_only=False)
    try:
        if workbook.sheetnames != [SHEET]:
            raise ValueError(f"Unexpected workbook sheets: {workbook.sheetnames}")
        sheet = workbook[SHEET]
        iterator = sheet.iter_rows()
        if [cell.value for cell in next(iterator)] != HEADERS:
            raise ValueError("Workbook headers do not match rev03")
        def checked_rows():
            for number, cells in enumerate(iterator, 2):
                for cell in cells:
                    if cell.data_type in ("f", "e"):
                        raise ValueError(f"Formula/error at {SHEET}!{cell.coordinate}; do not use cached or coerced values")
                yield number, [cell.value for cell in cells]
        unique, raw_locations, lineage, quality = normalize_rows(checked_rows())
        quality["workbook"] = {"sheets": [{"name": SHEET, "state": sheet.sheet_state,
                                           "rows_with_header": sheet.max_row, "columns": sheet.max_column}],
                               "date_encoding": "Gregorian Year and numeric Month columns; no day supplied",
                               "formulas": 0, "cell_errors": 0}
    finally:
        workbook.close()
    # Source IDs are workbook identifiers, not inferred administrative codes.
    locations = [{"source_id": source_id, "source_tambon_en": row[0], "source_amphoe_en": row[1],
                  "irrigation_status_raw": row[2], "irrigation_status": irrigation_status(row[2]),
                  "irrigation_label_th": IRRIGATION_LABELS[irrigation_status(row[2])]}
                 for source_id, row in sorted(raw_locations.items(), key=lambda item: int(item[0]))]
    labels = defaultdict(list)
    for source_id, row in raw_locations.items():
        labels[(row[1], row[0])].append(source_id)
    values, packed, run_counts = [], defaultdict(dict), defaultdict(Counter)
    for (source_month, source_id), row in sorted(unique.items(), key=lambda item: (item[0][0], int(item[0][1]))):
        packed[source_month][source_id] = row["risks"]
        for horizon, risk in enumerate(row["risks"], 1):
            values.append({"source_id": source_id, "source_month": source_month,
                           "horizon_months": horizon, "target_month": target_month(source_month, horizon),
                           "risk_value": risk, "scope_status": "out_of_scope" if risk is None else "in_scope",
                           "source_first_row": row["source_first_row"], "source_row_count": row["source_row_count"]})
            run_counts[(source_month, horizon)]["null" if risk is None else str(risk)] += 1
    runs = [{"source_month": month, "horizon_months": horizon, "target_month": target_month(month, horizon),
             "record_count": sum(counts.values()), "in_scope_count": sum(counts[str(r)] for r in range(3)),
             "out_of_scope_count": counts["null"], "no_signal_count": counts["0"],
             "moderate_count": counts["1"], "high_count": counts["2"]}
            for (month, horizon), counts in sorted(run_counts.items())]
    quality.update({"forecast_value_count": len(values), "forecast_run_count": len(runs),
                    "normalized_risk_counts": dict(Counter("null" if v["risk_value"] is None else str(v["risk_value"]) for v in values)),
                    "irrigation_location_counts": dict(Counter(item["irrigation_status"] for item in locations)),
                    "target_month_start": min(row["target_month"] for row in runs),
                    "target_month_end": max(row["target_month"] for row in runs),
                    "missing_location_months": 0, "missing_horizons": 0,
                    "source_district_label_count": len({r[1] for r in raw_locations.values()}),
                    "duplicate_location_labels": [{"amphoe_en": label[0], "tambon_en": label[1], "source_ids": ids}
                                                  for label, ids in sorted(labels.items()) if len(ids) > 1],
                    "latest_source_month_runs": runs[-6:]})
    meta = {"schema_version": 3, "dataset_id": "drought-rev03-" + source_hash[:12],
            "source_workbook": workbook_path.name, "source_sha256": source_hash, "source_sheet": SHEET,
            "source_month_role": "forecast_origin_month", "target_month_rule": "source_month + horizon_months",
            "time_precision": "month", "horizons": [1, 2, 3, 4, 5, 6], "risk_labels_th": RISK_LABELS,
            "irrigation_labels_th": IRRIGATION_LABELS, "blank_risk": "out_of_scope", "missing_record": "missing_record_not_out_of_scope",
            "forecast_primary_key": ["dataset_id", "source_id", "source_month", "horizon_months"],
            "location_identity": "source_id from workbook; no administrative-code mapping or external enrichment",
            "publication_status": "normalized_only_not_published"}
    archive = {"meta": meta, "locations": locations, "sourceMonths": sorted(packed),
               "runs": runs, "packedRiskBySourceMonth": dict(packed)}
    if source_hash != sha256(workbook_path):
        raise ValueError("Source workbook changed during normalization")
    output.mkdir(parents=True)
    write_csv(output / "forecast_values.csv", values, list(values[0]))
    write_csv(output / "locations.csv", locations, list(locations[0]))
    write_csv(output / "forecast_runs.csv", runs, list(runs[0]))
    write_csv(output / "source_row_lineage.csv", lineage, list(lineage[0]))
    write_json(output / "forecast_archive.json", archive)
    write_json(output / "data_quality.json", quality)
    manifest = {**meta, "csv_encoding": "UTF-8 with BOM", "csv_null_encoding": "empty field; scope_status remains explicit",
                "location_primary_key": ["source_id"], "run_primary_key": ["source_month", "horizon_months"],
                "location_join": "forecast_values.source_id = locations.source_id; source IDs are text identifiers",
                "provenance_join": "source_row_lineage.retained_source_row = forecast_values.source_first_row",
                "irrigation_time_scope": "one stable workbook attribute per location; no independent observation date supplied",
                "normalizer_sha256": sha256(Path(__file__)),
                "files": {path.name: {"sha256": sha256(path), "bytes": path.stat().st_size} for path in sorted(output.iterdir())}}
    write_json(output / "manifest.json", manifest)
    return {"output": str(output), "quality": quality}


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--output-dir", type=Path, required=True)
    args = parser.parse_args()
    print(json.dumps(build(args.workbook, args.output_dir), ensure_ascii=False, indent=2))
